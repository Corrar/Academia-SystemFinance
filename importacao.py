"""Cadastro rápido: lê fichas coladas do Excel/Google Planilhas, CSV ou XLSX e valida antes de salvar.

As linhas viram dicionários com os campos do aluno; `validar` aponta os erros de cada
linha (para o admin corrigir na tela de revisão) e evita cadastros duplicados.
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta

import db

CAMPOS = ("nome", "telefone", "email", "cpf", "plano", "data_vencimento", "data_inscricao")
MAX_LINHAS = 500

# nome da coluna na planilha (sem acento, minúsculo) -> campo
CABECALHOS = {
    "nome": "nome", "nome completo": "nome", "aluno": "nome", "nome do aluno": "nome", "cliente": "nome",
    "telefone": "telefone", "whatsapp": "telefone", "whats": "telefone", "celular": "telefone", "fone": "telefone",
    "tel": "telefone", "contato": "telefone",
    "email": "email", "e-mail": "email", "e mail": "email",
    "cpf": "cpf",
    "plano": "plano", "modalidade": "plano",
    "vencimento": "data_vencimento", "validade": "data_vencimento", "vence": "data_vencimento",
    "valido ate": "data_vencimento", "proximo vencimento": "data_vencimento", "data de vencimento": "data_vencimento",
    "pago ate": "data_vencimento",
    "inscricao": "data_inscricao", "matricula": "data_inscricao", "data de inscricao": "data_inscricao",
    "data da matricula": "data_inscricao", "cadastro": "data_inscricao", "data de cadastro": "data_inscricao",
}
# sem cabeçalho: ordem mais comum numa ficha
ORDEM_PADRAO = ("nome", "telefone", "email", "plano", "data_vencimento")


def _normalizar(texto) -> str:
    texto = unicodedata.normalize("NFD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto.strip().lower().replace("_", " ").replace(":", ""))


def ler_data(valor) -> str:
    """Aceita 25/12/2026, 25/12/26, 2026-12-25, 25-12-2026, 25.12.2026, datas do Excel. Retorna AAAA-MM-DD ou ''."""
    if valor in (None, ""):
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, (int, float)) and 20000 < valor < 80000:  # número de série do Excel
        return (date(1899, 12, 30) + timedelta(days=int(valor))).isoformat()
    texto = str(valor).strip()
    if re.fullmatch(r"\d{5}(\.0)?", texto):
        return ler_data(float(texto))
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", texto)
    if m:
        a, me, d = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})", texto)
        if not m:
            return ""
        d, me, a = (int(x) for x in m.groups())
        if a < 100:
            a += 2000
    try:
        return date(a, me, d).isoformat()
    except ValueError:
        return ""


def formatar_telefone(valor) -> str:
    """(11) 99999-0000 quando dá para reconhecer o número; senão devolve como veio."""
    texto = str(valor or "").strip()
    if texto.endswith(".0"):  # número lido do Excel como float
        texto = texto[:-2]
    digitos = re.sub(r"\D", "", texto)
    if digitos.startswith("55") and len(digitos) in (12, 13):
        digitos = digitos[2:]
    if len(digitos) == 11:
        return f"({digitos[:2]}) {digitos[2:7]}-{digitos[7:]}"
    if len(digitos) == 10:
        return f"({digitos[:2]}) {digitos[2:6]}-{digitos[6:]}"
    return texto


def _achar_plano(texto, planos) -> str:
    """Id do plano pelo nome ("mensal") ou pelo valor ("120" / "R$ 120,00"); '' se não achar."""
    alvo = _normalizar(texto)
    if not alvo:
        return ""
    for p in planos:
        if _normalizar(p["nome"]) == alvo:
            return str(p["id"])
    for p in planos:
        if alvo in _normalizar(p["nome"]) or _normalizar(p["nome"]) in alvo:
            return str(p["id"])
    numero = re.sub(r"[^\d,\.]", "", alvo)
    if numero:
        try:
            valor = float(numero.replace(".", "").replace(",", ".")) if "," in numero else float(numero)
        except ValueError:
            return ""
        for p in planos:
            if abs(p["valor"] - valor) < 0.01:
                return str(p["id"])
    return ""


def _de_tabela(tabela: list[list], planos) -> list[dict]:
    tabela = [[("" if c is None else c) for c in linha] for linha in tabela if any(str(c or "").strip() for c in linha)]
    if not tabela:
        return []
    cabecalho = [CABECALHOS.get(_normalizar(c)) for c in tabela[0]]
    if "nome" in cabecalho:
        colunas, dados = cabecalho, tabela[1:]
    else:
        colunas, dados = list(ORDEM_PADRAO), tabela
    linhas = []
    for bruta in dados[:MAX_LINHAS]:
        linha = {campo: "" for campo in CAMPOS}
        for campo, valor in zip(colunas, bruta):
            if campo and not linha[campo]:
                linha[campo] = valor
        linhas.append({
            "nome": " ".join(str(linha["nome"]).split()),
            "telefone": formatar_telefone(linha["telefone"]),
            "email": str(linha["email"]).strip().lower(),
            "cpf": str(linha["cpf"]).strip(),
            "plano_id": _achar_plano(linha["plano"], planos),
            "plano_lido": str(linha["plano"]).strip(),
            "data_vencimento": ler_data(linha["data_vencimento"]),
            "venc_lido": str(linha["data_vencimento"]).strip(),
            "data_inscricao": ler_data(linha["data_inscricao"]),
        })
    return linhas


def ler_texto(texto: str, planos) -> list[dict]:
    """Texto colado do Excel/Google Planilhas (colunas separadas por tab) ou CSV com ; ou ,."""
    texto = texto.strip("﻿\n\r ")
    if not texto:
        return []
    amostra = texto[:2000]
    separador = "\t" if "\t" in amostra else (";" if amostra.count(";") >= amostra.count(",") else ",")
    return _de_tabela(list(csv.reader(io.StringIO(texto), delimiter=separador)), planos)


def ler_arquivo(nome: str, conteudo: bytes, planos) -> list[dict]:
    """Arquivo .csv (UTF-8 ou o padrão do Excel no Brasil) ou .xlsx."""
    if nome.lower().endswith(".xlsx"):
        from openpyxl import load_workbook  # só carrega quando precisa

        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
        return _de_tabela([list(r) for r in planilha.iter_rows(values_only=True)], planos)
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return ler_texto(conteudo.decode(codificacao), planos)
        except UnicodeDecodeError:
            continue
    return []


def linhas_do_formulario(form) -> list[dict]:
    """Fichas digitadas na tela (campos com [] repetidos por linha)."""
    colunas = {campo: form.getlist(campo) for campo in ("nome", "telefone", "email", "cpf", "plano_id",
                                                         "data_vencimento", "data_inscricao")}
    total = max((len(v) for v in colunas.values()), default=0)
    linhas = []
    for i in range(total):
        linha = {campo: (valores[i].strip() if i < len(valores) else "") for campo, valores in colunas.items()}
        linha["telefone"] = formatar_telefone(linha["telefone"])
        linhas.append(linha)
    return linhas


def preenchida(linha: dict) -> bool:
    return any(linha.get(c) for c in ("nome", "telefone", "email", "cpf", "data_vencimento"))


def validar(linhas: list[dict]) -> dict[int, list[str]]:
    """Erros por índice de linha. Linhas totalmente vazias são ignoradas."""
    existentes = db.contatos_existentes()
    vistos = {"telefone": {}, "email": {}, "cpf": {}}
    erros: dict[int, list[str]] = {}
    for i, linha in enumerate(linhas):
        if not preenchida(linha):
            continue
        problemas = []
        if not linha["nome"]:
            problemas.append("falta o nome")
        if not ler_data(linha.get("data_vencimento")):
            problemas.append("informe até quando o plano vale (vencimento)")
        if linha.get("data_inscricao") and not ler_data(linha["data_inscricao"]):
            problemas.append("data de inscrição inválida")
        if linha["email"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", linha["email"]):
            problemas.append("e-mail inválido")
        telefone = re.sub(r"\D", "", linha["telefone"])[-11:]
        if linha["telefone"] and len(telefone) < 10:
            problemas.append("telefone sem DDD")
        chaves = {"telefone": telefone if len(telefone) >= 10 else "", "email": linha["email"].lower(),
                  "cpf": re.sub(r"\D", "", linha["cpf"])}
        for campo, valor in chaves.items():
            if not valor:
                continue
            rotulo = {"telefone": "telefone", "email": "e-mail", "cpf": "CPF"}[campo]
            if valor in existentes[campo]:
                problemas.append(f"{rotulo} já cadastrado para {existentes[campo][valor]}")
            elif valor in vistos[campo]:
                problemas.append(f"{rotulo} repetido na ficha {vistos[campo][valor] + 1}")
            else:
                vistos[campo][valor] = i
        if problemas:
            erros[i] = problemas
    return erros


def modelo_csv() -> str:
    """Planilha modelo (separada por ponto e vírgula, que o Excel no Brasil abre direto)."""
    hoje = db.hoje()
    exemplos = [
        ["Nome", "WhatsApp", "E-mail", "CPF", "Plano", "Vencimento", "Inscrição"],
        ["Maria Souza", "(11) 98888-7777", "maria@email.com", "", "Mensal",
         (hoje + timedelta(days=12)).strftime("%d/%m/%Y"), (hoje - timedelta(days=200)).strftime("%d/%m/%Y")],
        ["João Lima", "(11) 97777-6666", "", "", "Trimestral",
         (hoje + timedelta(days=40)).strftime("%d/%m/%Y"), (hoje - timedelta(days=50)).strftime("%d/%m/%Y")],
    ]
    saida = io.StringIO()
    csv.writer(saida, delimiter=";", lineterminator="\r\n").writerows(exemplos)
    return "﻿" + saida.getvalue()
